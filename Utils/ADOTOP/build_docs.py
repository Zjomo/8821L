from pathlib import Path
from copy import deepcopy
import shutil
import cv2
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
REF_DIR = Path(r'C:\Users\Mr\AppData\Local\Temp\adotop_docx_inspect')
OUT = ROOT / 'generated_docs'
OUT.mkdir(exist_ok=True)

def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color='D9D9D9', sz='6'):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None: borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:'+edge; el = borders.find(qn(tag))
        if el is None: el = OxmlElement(tag); borders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),sz); el.set(qn('w:color'),color)

def add_table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=str(h); set_cell_shading(c,'1F4E78'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.color.rgb=RGBColor(255,255,255); r.bold=True; r.font.size=Pt(9)
        set_cell_border(c)
    trPr = t.rows[0]._tr.get_or_add_trPr(); hdr = OxmlElement('w:tblHeader'); hdr.set(qn('w:val'), 'true'); trPr.append(hdr)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_border(cells[i])
            if ri%2==1: set_cell_shading(cells[i],'F2F6FA')
            for p in cells[i].paragraphs:
                for r in p.runs: r.font.size=Pt(9)
    return t

def setup_doc(ref, title):
    d=Document(str(ref)); clear_body(d)
    sec=d.sections[0]; sec.top_margin=Inches(0.8); sec.bottom_margin=Inches(0.8); sec.left_margin=Inches(0.9); sec.right_margin=Inches(0.9)
    p=d.add_paragraph(); p.style='Title'; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(title); r.font.name='黑体'; r.font.size=Pt(22); r.bold=True
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('项目文档'); r.font.name='黑体'; r.font.size=Pt(14)
    d.add_paragraph('')
    return d

def h(d, text, level=1):
    p=d.add_paragraph(text, style=f'Heading {level}'); p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(5); return p
def para(d, text):
    p=d.add_paragraph(text); p.paragraph_format.first_line_indent=Inches(0.3); p.paragraph_format.line_spacing=1.25; p.paragraph_format.space_after=Pt(5); return p
def caption(d, text):
    p=d.add_paragraph(text); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(8); 
    for r in p.runs: r.italic=True; r.font.size=Pt(9)

def add_image(d, path, width=5.9):
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; run=p.add_run(); run.add_picture(str(path), width=Inches(width))
    for inline in run._r.xpath('.//wp:docPr'):
        inline.set('descr', path.stem); inline.set('title', path.stem)
    return p

def make_figures():
    figdir=OUT/'figures'; figdir.mkdir(exist_ok=True)
    frame=next((ROOT/'Dataset'/'VideoFrames').glob('*.jpg'))
    im=cv2.imread(str(frame));
    # foreground/edge illustration
    gray=cv2.cvtColor(im,cv2.COLOR_BGR2GRAY); _,mask=cv2.threshold(gray,0,255,cv2.THRESH_BINARY+cv2.THRESH_OTSU)
    fg=cv2.bitwise_and(im,im,mask=mask); tri=np.hstack([im,fg,cv2.cvtColor(mask,cv2.COLOR_GRAY2BGR)])
    cv2.imwrite(str(figdir/'foreground_pipeline.jpg'),tri)
    edges=cv2.Canny(gray,50,150); lines=cv2.HoughLinesP(edges,1,np.pi/180,50,minLineLength=30,maxLineGap=10)
    out=im.copy(); best=None; bl=0
    if lines is not None:
        for l in lines[:,0]:
            x1,y1,x2,y2=map(int,l); le=((x2-x1)**2+(y2-y1)**2)**0.5
            if le>bl: bl=le; best=(x1,y1,x2,y2)
    angle=0
    if best:
        x1,y1,x2,y2=best; angle=abs(np.degrees(np.arctan2(y2-y1,x2-x1))); angle=180-angle if angle>90 else angle; cv2.line(out,(x1,y1),(x2,y2),(0,0,255),3); cv2.putText(out,f'angle={angle:.1f} deg',(20,45),cv2.FONT_HERSHEY_SIMPLEX,1.0,(0,255,255),2)
    cv2.imwrite(str(figdir/'edge_angle.jpg'),out)
    return figdir, frame

def build_seg():
    fig,frame=make_figures(); d=setup_doc(REF_DIR/'ref1.docx','基于YOLOv8实例分割与前后景融合的微纳物体视觉分析系统')
    h(d,'目录',1)
    for x in ['1 项目概况','1.1 背景和基础','1.2 场景和价值','1.3 所需支持','2 项目规划','2.1 整体目标','2.2 技术创新点','2.3 硬件设备','2.4 算法原理分析','3 实施方案','3.1 技术可行性分析','3.2 技术细节','3.3 系统测试与精度验证','4 参考资料']:
        p=d.add_paragraph(x); p.paragraph_format.left_indent=Inches(0.25)
    h(d,'项目概况',1); h(d,'背景和基础',2)
    para(d,'微纳加工、显微观测和实验教学中，目标物体常以视频流或屏幕采集画面的形式出现。传统人工观察依赖操作者经验，难以同时完成前景提取、目标实例分割、几何量测和结果留存。ADOTOP项目将视频前后景分离、YOLOv8分割、边缘与轮廓分析整合到统一桌面工具中，形成可重复的视觉分析流程。')
    para(d,'代码仓库由Main.py统一入口、ForegroundAndBackground前后景模块、YoloUI数据准备/训练/推理模块以及EdgeDetect几何识别模块组成。现有实现支持视频文件、静态图像和屏幕ROI三类输入，能够输出原图、分割结果、二值掩码、检测框和JSON快照。')
    add_image(d,fig/'foreground_pipeline.jpg',5.8); caption(d,'图1.1 前后景提取与二值掩码处理示意')
    h(d,'场景和价值',2)
    para(d,'系统适用于显微镜视频中的物体分离、材料实验中的目标轮廓提取、屏幕内局部区域的在线监测以及分割模型训练前后的快速核验。其价值不在于替代实验人员，而在于将输入、参数、结果和快照保存为可追溯记录，降低重复观察与手工记录成本。')
    h(d,'所需支持',2); h(d,'硬件平台的论证和选择',3)
    add_table(d,['设备','代码中的用途','当前边界'],[['USB/屏幕采集设备','OpenCV VideoCapture、Pillow ImageGrab、pyautogui ROI','依赖本机驱动与桌面权限'],['GPU或CPU工作站','Ultralytics YOLOv8训练与推理','训练参数由args.yaml记录，CPU模式可运行但速度较慢'],['显示器与输入设备','PyQt6 ROI选择、播放、暂停、拖动定位','属于桌面交互，不等同于自动化机械执行']])
    h(d,'软件平台的论证和选择',3)
    add_table(d,['软件层','仓库证据'],[['界面','PyQt6，Main.py与YoloUI/yolo_seg_gui.py'],['视觉处理','OpenCV、NumPy，Canny、HoughLinesP、阈值分割、轮廓分析'],['深度学习','Ultralytics YOLOv8-seg，支持数据准备、训练、图片/视频/ROI推理'],['数据与结果','YAML数据集配置、JSON快照、runs/segment训练结果']])
    h(d,'项目规划',1); h(d,'整体目标',2)
    para(d,'系统采用感知、分析、交互三层结构：感知层从视频、图像或屏幕ROI获得帧；分析层完成前后景分离、实例分割、轮廓分类、最长直线检测与角度计算；交互层提供参数配置、播放控制、ROI选择、手动标注、训练和评估入口。')
    h(d,'技术创新点',2); h(d,'作品难点',3)
    para(d,'1. 多输入源统一处理：视频、图片和屏幕ROI共享同一套显示与分析逻辑，同时处理暂停、拖动定位、循环播放和异常输入。')
    para(d,'2. 传统视觉与深度模型协同：前景模块提供bright、otsu、blue、yellow等可解释阈值方法；YoloUI提供YOLOv8实例分割，EdgeDetect再对掩码和轮廓进行几何分析。')
    para(d,'3. 结果可追溯：推理结果可保存为PNG与JSON，训练目录保留args.yaml、results.csv、混淆矩阵和PR/F1曲线。')
    h(d,'作品创新点',3)
    para(d,'系统将可解释的颜色/灰度分割、学习型实例分割和几何规则融合在一个可操作界面中；StableLineDetector使用指数滑动平均和跳变阈值抑制角度抖动；手动模板标注可反哺轮廓分类。')
    h(d,'主要硬件设备',2)
    para(d,'当前仓库未绑定固定工业相机、运动平台或激光执行器。输入侧以OpenCV/Pillow/pyautogui抽象，输出侧为屏幕显示、图片和JSON记录。因此本项目定位为视觉分析与实验辅助软件，不宣称具备机械切割或闭环执行能力。')
    h(d,'算法原理分析',2); h(d,'前后景分离模块',3)
    para(d,'segment_foreground_background根据方法选择HSV颜色阈值或Otsu灰度阈值，再通过形态学清理和最小连通域面积过滤得到二值掩码。GUI同时展示原图、前景提取结果和二值图，便于调整方法与min_area参数。')
    h(d,'YOLOv8实例分割模块',3)
    para(d,'dataset_utils.py支持YOLO检测、YOLO分割和COCO格式解析，并可执行数据集划分与YAML生成。训练脚本以Ultralytics YOLOv8-seg为核心，训练目录保留模型参数和评估图。仓库中可见一次yolov8m-seg、12 epoch、640输入、batch 4、CPU训练记录；results.csv中的部分指标为0，表明该记录不能直接作为稳定精度结论。')
    h(d,'边缘与角度检测模块',3)
    para(d,'EdgeDetect使用灰度预处理、阈值与轮廓提取完成圆底、长条和旋转体的分类；对最长直线使用Canny边缘和HoughLinesP检测，并将atan2结果归一化到0至90度。StableLineDetector对角度变化超过5度的跳变进行拒绝，并以EMA平滑连续帧。')
    add_image(d,fig/'edge_angle.jpg',5.8); caption(d,'图2.1 Canny与HoughLinesP最长边角度示意')
    h(d,'实施方案',1); h(d,'技术可行性分析',2)
    add_table(d,['环节','已实现能力','验证方式'],[['输入管理','视频/图片/屏幕ROI，播放、暂停、seek','Main.py、video_seg_ui.py'],['前景分离','4种阈值策略与连通域清理','video_seg.py函数级调用'],['实例分割','数据格式转换、训练、图片/视频/ROI推理','YoloUI目录与runs结果'],['几何分析','轮廓分类、模板相似度、最长边角度','EdgeDetect模块']])
    h(d,'技术细节',2); h(d,'系统主程序设计',3)
    para(d,'Main.py通过UnifiedMainWindow组织ForegroundPage与EdgeDetectPage；后台线程负责视频和屏幕ROI采集，信号传递原图、分割图、二值图和角度。YoloUI独立提供数据集、训练、预测和评估标签页。')
    h(d,'系统测试与精度验证',2)
    para(d,'仓库包含真实帧样本、TIFF图像、JSON标注和训练可视化文件，可用于回放式验证。建议验收关注：ROI坐标是否正确、不同阈值方法是否产生可用掩码、视频seek后是否停在目标帧、模型推理失败时是否给出错误提示、快照PNG与JSON是否成对生成。当前材料未提供独立测试集上的完整mAP统计，故不将训练曲线或单次可视化结果表述为生产精度。')
    h(d,'参考资料',1)
    for x in ['[1] Ultralytics. YOLOv8 Documentation and Source Code.', '[2] Bradski G. The OpenCV Library. Dr. Dobb’s Journal of Software Tools, 2000.', '[3] ADOTOP仓库 Main.py 与 ForegroundAndBackground 模块.', '[4] ADOTOP仓库 YoloUI 数据准备、训练与推理模块.', '[5] ADOTOP仓库 EdgeDetect 识别与界面模块.']:
        para(d,x)
    d.save(OUT/'ADOTOP_基于YOLOv8实例分割与前后景融合的微纳物体视觉分析系统.docx')

def build_edge():
    fig,frame=make_figures(); d=setup_doc(REF_DIR/'ref2.docx','基于多模态视觉采集与鲁棒几何分析的微纳物体角度检测系统')
    h(d,'项目概况',1); h(d,'背景和基础',2)
    para(d,'微纳加工和光学实验中，目标物体的姿态、长边方向和相对角度是判断装配状态与实验重复性的基础信息。ADOTOP中的EdgeDetect模块面向视频、图片和屏幕ROI，提供阈值预处理、轮廓提取、模板分类、最长边检测和角度显示功能，并由PyQt6界面统一管理输入、播放、ROI与结果。')
    para(d,'系统的角度测量来自图像几何，不包含Newport控制器、QPD、EKF或四轴伺服执行链路。相关硬件未在本仓库中实现，因此文档将系统边界限定为视觉检测与实验辅助软件。')
    add_image(d,fig/'edge_angle.jpg',5.8); caption(d,'图1.1 最长直线检测与角度显示')
    h(d,'所需支持',2); h(d,'硬件平台的论证和选择',3)
    add_table(d,['设备','用途','状态'],[['摄像头或视频文件','提供待分析帧','OpenCV接口已实现'],['桌面屏幕','提供screen模式和ROI截取','Pillow/pyautogui依赖权限'],['CPU/GPU工作站','运行OpenCV与可选YOLO推理','CPU可执行传统视觉流程'],['输入设备','鼠标框选ROI、键盘控制','PyQt6/cv2交互']])
    h(d,'软件平台的论证和选择',3)
    add_table(d,['模块','实现'],[['界面','PyQt6 QMainWindow、QDockWidget、QTableWidget'],['图像处理','OpenCV Canny、阈值、findContours、minAreaRect、HoughLinesP'],['模型辅助','YoloUI中的Ultralytics YOLOv8-seg，可作为分割前端'],['持久化','JSON快照、模板JSON、训练runs目录']])
    h(d,'项目规划',1); h(d,'整体目标',2)
    para(d,'系统按采集、预处理、几何识别、结果展示四层组织。采集层支持媒体文件和屏幕ROI；预处理层完成灰度化、二值化、形态学清理与重叠区域拆分；几何识别层输出轮廓中心、面积、周长、边界框、类别和角度；展示层提供原图/分割图、状态表、进度条、快照和手动模板标注。')
    h(d,'技术创新点',2); h(d,'作品难点',3)
    para(d,'1. 面向不稳定视频流的状态管理：VideoThread与ScreenRoiThread分别管理暂停、停止、seek和参数更新，避免输入模式互相冲突。')
    para(d,'2. 规则与模板结合的对象识别：contour_metrics计算几何特征，contour_similarity比较模板形状，classify_contours输出circle、rod、rotator等类别。')
    para(d,'3. 长边角度的稳健输出：HoughLinesP选取最长线段，StableLineDetector用跳变阈值与EMA抑制瞬时误检。')
    h(d,'作品创新点',3)
    para(d,'系统把实时预览、ROI操作、几何量测、模板维护和JSON证据留存放入同一工作台，能够在更换材料和输入源后快速复核参数；同时保留YOLO分割入口，为后续从规则视觉过渡到学习型分割提供兼容接口。')
    h(d,'主要硬件设备',2)
    para(d,'仓库没有四轴压电反射镜、QPD、波前传感器或激光器控制代码。角度检测结果仅代表图像平面中的线段方向，不能直接解释为光束角度、机械轴角度或物理世界绝对姿态。')
    h(d,'算法原理分析',2); h(d,'图像预处理与轮廓提取',3)
    para(d,'preprocess_frame将输入转为灰度并进行阈值处理，split_overlapped_mask用于处理重叠区域，extract_contours提取候选轮廓。contour_metrics计算面积、周长、中心和边界框，为类别判定与结果表提供统一数据结构。')
    h(d,'模板相似度与类别判定',3)
    para(d,'ManualAnnotator允许用户在当前帧上标注模板并保存。load_templates读取JSON模板，contour_similarity对候选轮廓和模板进行形状比较，classify_contours综合几何规则与模板分数输出Detection。')
    h(d,'最长边与角度检测',3)
    para(d,'detect_longest_line_single对灰度图执行Canny和HoughLinesP，遍历候选线段并选取长度最大的线段；角度由atan2计算后折算到0至90度。StableLineDetector在连续帧中拒绝超过5度的突变，并以alpha=0.3进行指数平滑。')
    h(d,'实施方案',1); h(d,'技术细节',2); h(d,'系统主程序设计',3)
    para(d,'UnifiedMainWindow以标签页组织前后景页面和EdgeDetect页面。EdgeDetect页面提供媒体打开、屏幕模式、ROI、手动标注、播放/暂停、录制和快照动作；识别结果在右侧表格展示中心、角度、置信度、面积和周长。')
    h(d,'系统测试与精度验证',2)
    add_table(d,['测试项','操作','验收关注点'],[['媒体输入','加载jpg、tif或视频并播放','首帧可读、循环播放、异常路径提示'],['ROI模式','屏幕截图后鼠标框选区域','坐标映射正确、ROI越界不崩溃'],['角度检测','输入含明显长边的帧','最长线段与角度叠加显示，跳变被抑制'],['结果留存','执行快照与手动标注','PNG/JSON生成，模板可再次加载']])
    para(d,'现有仓库包含大量VideoFrames样本、TIFF图像、JSON标注和YOLO训练产物，但未提供带物理角度真值的独立标定集。因此可报告算法输出和流程稳定性，不应将像素角度直接宣称为微弧度级物理测量精度。')
    h(d,'参考资料',1)
    for x in ['[1] Bradski G. The OpenCV Library. Dr. Dobb’s Journal of Software Tools, 2000.', '[2] Ultralytics. YOLOv8 Documentation and Source Code.', '[3] ADOTOP仓库 EdgeDetect/recognition_module.py.', '[4] ADOTOP仓库 EdgeDetect/three_object_detector.py.', '[5] ADOTOP仓库 Main.py 与 YoloUI/yolo_seg_gui.py.']:
        para(d,x)
    d.save(OUT/'ADOTOP_基于多模态视觉采集与鲁棒几何分析的微纳物体角度检测系统.docx')

if __name__=='__main__':
    build_seg(); build_edge(); print('generated', *OUT.glob('*.docx'), sep='\n')
