"""
验证“6. 光谱仪通信 / 单次光谱采集”下方的
“保存当前光谱数据（背景光）”按钮及保存逻辑。
"""

from __future__ import annotations

import importlib.util
import inspect
import sys
import tempfile
from pathlib import Path
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook

AUTOZOOM_ROOT = Path(__file__).resolve().parent
if str(AUTOZOOM_ROOT) not in sys.path:
    sys.path.insert(0, str(AUTOZOOM_ROOT))


def _import_7_25():
    """动态导入 7_25 主模块。"""
    module_name = "measurement_workflow_7_25"
    if module_name in sys.modules:
        return sys.modules[module_name]

    module_path = AUTOZOOM_ROOT / "0_measurement_workflow_real_virtual_same_detection_7_25.py"
    spec = importlib.util.spec_from_file_location(module_name, str(module_path))
    module = importlib.util.module_from_spec(spec)
    sys.modules[module_name] = module
    spec.loader.exec_module(module)
    return module


def _get_main_source() -> str:
    module_path = AUTOZOOM_ROOT / "0_measurement_workflow_real_virtual_same_detection_7_25.py"
    return module_path.read_text(encoding="utf-8")


def test_ui_contains_background_light_button() -> None:
    """主文件源码中应包含新按钮及其命令绑定。"""
    source = _get_main_source()

    assert "保存当前光谱数据（背景光）" in source
    assert "save_current_spectrum_background_light_thread" in source
    print("PASS: ui_contains_background_light_button")


def test_save_single_spectrum_to_xlsx_accepts_save_dir_and_tag() -> None:
    """save_single_spectrum_to_xlsx 应支持自定义 save_dir 与 tag。"""
    _module = _import_7_25()
    MeasurementConfig = _module.MeasurementConfig
    MeasurementWorkflow = _module.MeasurementWorkflow

    with tempfile.TemporaryDirectory() as tmpdir:
        cfg = MeasurementConfig()
        cfg.save_root = tmpdir
        wf = MeasurementWorkflow(cfg)

        wf.context["raw_values"] = [1.0, 2.0, 3.0]
        wf.context["raw_filtered_values"] = [1.1, 2.1, 3.1]
        wf.context["raw_median_values"] = [1.2, 2.2, 3.2]
        wf.context["fit_values"] = [1.3, 2.3, 3.3]
        wf.context["x_axis_values"] = [400.0, 500.0, 600.0]

        save_dir = Path(tmpdir) / "custom_save"
        saved_path = wf.save_single_spectrum_to_xlsx(
            save_dir=save_dir,
            tag="background_light",
        )

        assert saved_path is not None
        assert saved_path.exists()
        assert saved_path.parent == save_dir
        assert "background_light" in saved_path.name

        wb = load_workbook(saved_path)
        ws = wb.active
        assert ws.title == "光谱数据"
        assert [cell.value for cell in ws[1]] == [
            "波长/索引",
            "原始强度",
            "阈值滤波",
            "中值滤波",
            "拟合曲线",
        ]
        assert ws.max_row == 4  # 1 表头 + 3 行数据

        print("PASS: save_single_spectrum_to_xlsx_accepts_save_dir_and_tag")


def test_save_single_spectrum_to_xlsx_default_uses_output_root_save() -> None:
    """默认参数下，保存路径应位于 output_root/save/{MM.DD} 且不含 tag。"""
    _module = _import_7_25()
    MeasurementConfig = _module.MeasurementConfig
    MeasurementWorkflow = _module.MeasurementWorkflow

    with tempfile.TemporaryDirectory() as tmpdir:
        cfg = MeasurementConfig()
        cfg.save_root = tmpdir
        wf = MeasurementWorkflow(cfg)

        wf.context["raw_values"] = [10.0, 20.0]
        wf.context["x_axis_values"] = [100.0, 200.0]

        saved_path = wf.save_single_spectrum_to_xlsx()

        assert saved_path is not None
        assert saved_path.exists()
        # 默认路径在 output_root / save 下
        assert str(saved_path).startswith(str(Path(tmpdir) / "save"))
        assert "background_light" not in saved_path.name
        assert saved_path.name.startswith("measurement_summary_")

        print("PASS: save_single_spectrum_to_xlsx_default_uses_output_root_save")


def test_save_current_spectrum_background_light_calls_workflow_save() -> None:
    """GUI 方法应调用 workflow 的保存方法，并传入 ./save 目录和 background_light tag。"""
    _module = _import_7_25()
    MeasurementWorkflowGUI = _module.MeasurementWorkflowGUI

    source = inspect.getsource(
        _module.MeasurementWorkflowGUI.save_current_spectrum_background_light
    )
    assert "save_dir = Path(\"./save\")" in source or 'save_dir = Path("./save")' in source
    assert 'tag="background_light"' in source

    with patch.object(MeasurementWorkflowGUI, "__init__", lambda self, root: None):
        gui = MeasurementWorkflowGUI(None)
        gui.tcp_status_var = MagicMock()

        fake_wf = MagicMock()
        fake_wf.save_single_spectrum_to_xlsx.return_value = Path(
            "./save/01.01/measurement_summary_background_light_20260101_120000.xlsx"
        )

        with patch.object(gui, "set_var", MagicMock()) as mock_set_var:
            with patch.object(gui, "ensure_workflow", return_value=fake_wf):
                with patch.object(gui, "log", MagicMock()):
                    gui.save_current_spectrum_background_light()

            call_kwargs = fake_wf.save_single_spectrum_to_xlsx.call_args.kwargs
            save_dir = call_kwargs.get("save_dir")
            assert "save" in save_dir.parts
            assert call_kwargs.get("tag") == "background_light"
            assert mock_set_var.called

        print("PASS: save_current_spectrum_background_light_calls_workflow_save")


def test_save_current_spectrum_background_light_no_data() -> None:
    """无光谱数据时，GUI 方法应更新状态为保存失败。"""
    _module = _import_7_25()
    MeasurementWorkflowGUI = _module.MeasurementWorkflowGUI

    with patch.object(MeasurementWorkflowGUI, "__init__", lambda self, root: None):
        gui = MeasurementWorkflowGUI(None)
        gui.tcp_status_var = MagicMock()

        fake_wf = MagicMock()
        fake_wf.save_single_spectrum_to_xlsx.return_value = None

        with patch.object(gui, "set_var", MagicMock()) as mock_set_var:
            with patch.object(gui, "ensure_workflow", return_value=fake_wf):
                with patch.object(gui, "log", MagicMock()):
                    gui.save_current_spectrum_background_light()

        status_call = mock_set_var.call_args[0][1]
        assert "保存失败" in status_call or "无可用光谱数据" in status_call

        print("PASS: save_current_spectrum_background_light_no_data")


if __name__ == "__main__":
    tests = [
        test_ui_contains_background_light_button,
        test_save_single_spectrum_to_xlsx_accepts_save_dir_and_tag,
        test_save_single_spectrum_to_xlsx_default_uses_output_root_save,
        test_save_current_spectrum_background_light_calls_workflow_save,
        test_save_current_spectrum_background_light_no_data,
    ]

    failed = 0
    for test in tests:
        try:
            test()
        except Exception as e:
            failed += 1
            print(f"FAIL: {test.__name__}: {e}")

    print(f"\n{len(tests) - failed}/{len(tests)} tests passed")
    if failed:
        sys.exit(1)
