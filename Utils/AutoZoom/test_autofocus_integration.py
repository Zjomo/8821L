"""
自动补焦参数集成与光谱数据保存测试。

覆盖：
  1. MeasurementConfig 默认补焦参数
  2. MeasurementWorkflowGUI 光谱补焦循环 UI 变量默认值与 _make_saf_config 同步
  3. MeasurementWorkflow 聚焦参考图建立
  4. MeasurementWorkflow FocusScore_ratio 计算与补焦触发判断
  5. MeasurementWorkflow.save_single_spectrum_to_xlsx 导出 xlsx
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path
from typing import Any, Dict

import cv2
import numpy as np
import pytest

AUTOZOOM_ROOT = Path(__file__).resolve().parent
if str(AUTOZOOM_ROOT) not in sys.path:
    sys.path.insert(0, str(AUTOZOOM_ROOT))

import importlib.util

_module_path = AUTOZOOM_ROOT / "0_measurement_workflow_real_virtual_same_detection_7_16.py"
_spec = importlib.util.spec_from_file_location(
    "measurement_workflow_7_16", str(_module_path)
)
_measurement_workflow_7_16 = importlib.util.module_from_spec(_spec)
sys.modules["measurement_workflow_7_16"] = _measurement_workflow_7_16
_spec.loader.exec_module(_measurement_workflow_7_16)  # type: ignore[union-attr]

MeasurementConfig = _measurement_workflow_7_16.MeasurementConfig
MeasurementWorkflow = _measurement_workflow_7_16.MeasurementWorkflow
MeasurementWorkflowGUI = _measurement_workflow_7_16.MeasurementWorkflowGUI


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest.fixture
def workflow() -> MeasurementWorkflow:
    """构造 virtual 模式 workflow，避免真实硬件。"""
    cfg = MeasurementConfig(
        hardware_mode="virtual",
        spectrometer_backend="mock",
        focus_roi=(100, 100, 200, 200),
    )
    wf = MeasurementWorkflow(cfg)
    return wf


@pytest.fixture
def sharp_image() -> np.ndarray:
    """生成清晰合成图（RGB，400x400）。"""
    rng = np.random.default_rng(42)
    img = rng.integers(0, 255, (400, 400, 3), dtype=np.uint8)
    # 添加锐利边缘提升高频分量
    cv2.rectangle(img, (120, 120), (280, 280), (255, 0, 0), 3)
    cv2.rectangle(img, (150, 150), (250, 250), (0, 255, 0), 3)
    return img


@pytest.fixture
def blurred_image(sharp_image: np.ndarray) -> np.ndarray:
    """对清晰图做高斯模糊，FocusScore 应显著下降。"""
    return cv2.GaussianBlur(sharp_image, (21, 21), 5)


# ---------------------------------------------------------------------------
# 需求 1：补焦参数与 UI
# ---------------------------------------------------------------------------


def test_measurement_config_focus_defaults() -> None:
    """MeasurementConfig 应包含完整补焦参数默认值。"""
    cfg = MeasurementConfig()
    assert cfg.focus_trigger_ratio == pytest.approx(0.95)
    assert cfg.focus_stop_ratio == pytest.approx(0.95)
    assert cfg.focus_trigger_count == 3
    assert cfg.focus_trigger_absolute is True
    assert cfg.focus_detection_only is False
    assert cfg.focus_z_enabled is True
    assert cfg.focus_z_axis == 1
    assert cfg.focus_z_speed == 100
    assert cfg.focus_z_accel == 100
    assert cfg.focus_search_strategy == "hill_climb"


def test_gui_saf_vars_defaults() -> None:
    """GUI 光谱补焦循环区域变量默认值应与 MeasurementConfig 一致。"""
    import tkinter as tk

    root = tk.Tk()
    try:
        gui = MeasurementWorkflowGUI(root)
        assert float(gui.saf_trigger_ratio_var.get()) == pytest.approx(0.95)
        assert float(gui.saf_stop_ratio_var.get()) == pytest.approx(0.95)
        assert int(gui.saf_trigger_count_var.get()) == 3
        assert bool(gui.saf_trigger_absolute_var.get()) is True
        assert bool(gui.saf_detection_only_var.get()) is False
        assert bool(gui.saf_passive_mode_var.get()) is True
        assert int(gui.saf_passive_attempts_var.get()) == 10
        assert int(gui.saf_passive_good_var.get()) == 5
        assert bool(gui.saf_disable_auto_stop_var.get()) is False
        assert bool(gui.saf_z_enabled_var.get()) is True
        assert int(gui.saf_z_axis_var.get()) == 1
        assert int(gui.saf_z_speed_var.get()) == 100
        assert int(gui.saf_z_accel_var.get()) == 100
        assert gui.saf_search_strategy_var.get() == "hill_climb"
        assert float(gui.saf_interval_var.get()) == pytest.approx(1.0)
        assert float(gui.saf_wait_between_spectrum_var.get()) == pytest.approx(120.0)
    finally:
        root.destroy()


def test_gui_make_saf_config() -> None:
    """_make_saf_config 应将 GUI 变量完整写入 AutofocusConfig。"""
    import tkinter as tk

    root = tk.Tk()
    try:
        gui = MeasurementWorkflowGUI(root)
        gui.saf_trigger_ratio_var.set(0.90)
        gui.saf_stop_ratio_var.set(0.92)
        gui.saf_trigger_count_var.set(2)
        gui.saf_trigger_absolute_var.set(False)
        gui.saf_detection_only_var.set(True)
        gui.saf_passive_mode_var.set(False)
        gui.saf_passive_attempts_var.set(5)
        gui.saf_passive_good_var.set(2)
        gui.saf_disable_auto_stop_var.set(True)
        gui.saf_z_enabled_var.set(False)
        gui.saf_z_axis_var.set(2)
        gui.saf_z_speed_var.set(50)
        gui.saf_z_accel_var.set(60)
        gui.saf_search_strategy_var.set("full_sweep")

        cfg = gui._make_saf_config()
        assert cfg.autofocus_focus_trigger_ratio == pytest.approx(0.90)
        assert cfg.autofocus_stop_ratio == pytest.approx(0.92)
        assert cfg.autofocus_focus_trigger_count == 2
        assert cfg.autofocus_trigger_absolute is False
        assert cfg.autofocus_detection_only is True
        assert cfg.autofocus_passive_mode is False
        assert cfg.autofocus_passive_max_attempts == 5
        assert cfg.autofocus_passive_consecutive_good == 2
        assert cfg.autofocus_passive_disable_auto_stop is True
        assert cfg.z_enabled is False
        assert cfg.z_axis == 2
        assert cfg.z_speed == 50
        assert cfg.z_accel == 60
        assert cfg.z_search_strategy == "full_sweep"
    finally:
        root.destroy()


# ---------------------------------------------------------------------------
# 需求 2：聚焦参考与补焦触发
# ---------------------------------------------------------------------------


def test_capture_focus_reference(workflow: MeasurementWorkflow, sharp_image: np.ndarray) -> None:
    """capture_focus_reference 应建立参考图并使 scorer 就绪。"""
    with tempfile.TemporaryDirectory() as tmpdir:
        workflow.output_root = Path(tmpdir)
        workflow._capture_current_focus_frame = lambda: sharp_image.copy()

        ok = workflow.capture_focus_reference(cycle_index=1)
        assert ok is True
        assert workflow._focus_reference_ready is True
        assert workflow._focus_reference_image is not None
        assert workflow._focus_scorer is not None
        assert workflow._focus_scorer.focus_reference_ready is True


def test_compute_focus_score_same_image(
    workflow: MeasurementWorkflow, sharp_image: np.ndarray
) -> None:
    """参考图与当前图相同，FocusScore_ratio 应接近 1.0。"""
    with tempfile.TemporaryDirectory() as tmpdir:
        workflow.output_root = Path(tmpdir)
        workflow._capture_current_focus_frame = lambda: sharp_image.copy()
        workflow.capture_focus_reference(cycle_index=1)

        score = workflow.compute_current_focus_score()
        assert score is not None
        assert 0.95 <= score <= 1.05


def test_compute_focus_score_blurred_image(
    workflow: MeasurementWorkflow, sharp_image: np.ndarray, blurred_image: np.ndarray
) -> None:
    """当前图模糊时，FocusScore_ratio 应低于触发阈值。"""
    with tempfile.TemporaryDirectory() as tmpdir:
        workflow.output_root = Path(tmpdir)
        workflow._capture_current_focus_frame = lambda: sharp_image.copy()
        workflow.capture_focus_reference(cycle_index=1)

        workflow._capture_current_focus_frame = lambda: blurred_image.copy()
        score = workflow.compute_current_focus_score()
        assert score is not None
        assert score < 0.95


def test_run_autofocus_if_needed_triggers(
    workflow: MeasurementWorkflow, sharp_image: np.ndarray, blurred_image: np.ndarray
) -> None:
    """连续低分应触发补焦；Z 轴禁用时仅返回 triggered=True 不抛异常。"""
    workflow.cfg.focus_z_enabled = False
    workflow.cfg.focus_trigger_count = 1

    with tempfile.TemporaryDirectory() as tmpdir:
        workflow.output_root = Path(tmpdir)
        workflow._capture_current_focus_frame = lambda: sharp_image.copy()
        workflow.capture_focus_reference(cycle_index=1)

        workflow._capture_current_focus_frame = lambda: blurred_image.copy()
        result = workflow.run_autofocus_if_needed(cycle_index=2)
        assert result["score"] is not None
        assert result["score"] < 0.95
        assert result["triggered"] is True


def test_run_autofocus_if_needed_no_trigger(
    workflow: MeasurementWorkflow, sharp_image: np.ndarray
) -> None:
    """分数达标时不触发补焦。"""
    workflow.cfg.focus_trigger_count = 1

    with tempfile.TemporaryDirectory() as tmpdir:
        workflow.output_root = Path(tmpdir)
        workflow._capture_current_focus_frame = lambda: sharp_image.copy()
        workflow.capture_focus_reference(cycle_index=1)

        score = workflow.run_autofocus_if_needed(cycle_index=2)
        assert score["score"] is not None
        assert score["triggered"] is False


def test_compute_focus_score_without_reference(workflow: MeasurementWorkflow) -> None:
    """未建立参考时，compute_current_focus_score 应返回 None 且不抛异常。"""
    score = workflow.compute_current_focus_score()
    assert score is None


# ---------------------------------------------------------------------------
# 需求 3：单次光谱数据保存
# ---------------------------------------------------------------------------


def test_save_single_spectrum_to_xlsx(workflow: MeasurementWorkflow) -> None:
    """save_single_spectrum_to_xlsx 应正确导出光谱数据到 xlsx。"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        pytest.skip("需要 openpyxl")

    with tempfile.TemporaryDirectory() as tmpdir:
        workflow.output_root = Path(tmpdir)
        n = 10
        workflow.context["raw_values"] = [float(i * 10) for i in range(n)]
        workflow.context["raw_filtered_values"] = [float(i * 10 + 1) for i in range(n)]
        workflow.context["raw_median_values"] = [float(i * 10 + 2) for i in range(n)]
        workflow.context["fit_values"] = [float(i * 10 + 3) for i in range(n)]
        workflow.context["x_axis_values"] = [float(400 + i) for i in range(n)]

        path = workflow.save_single_spectrum_to_xlsx()
        assert path is not None
        assert path.exists()
        assert path.suffix == ".xlsx"
        assert "save" in str(path)

        wb = load_workbook(path)
        ws = wb.active
        assert ws.title == "光谱数据"
        headers = [cell.value for cell in ws[1]]
        assert headers == ["波长/索引", "原始强度", "阈值滤波", "中值滤波", "拟合曲线"]
        assert ws.max_row == n + 1
        assert ws.max_column == 5

        first_row = [ws.cell(row=2, column=c).value for c in range(1, 6)]
        assert first_row[0] == 400.0
        assert first_row[1] == 0.0


def test_save_single_spectrum_to_xlsx_empty(workflow: MeasurementWorkflow) -> None:
    """无原始光谱数据时应跳过保存并返回 None。"""
    with tempfile.TemporaryDirectory() as tmpdir:
        workflow.output_root = Path(tmpdir)
        workflow.context["raw_values"] = []
        path = workflow.save_single_spectrum_to_xlsx()
        assert path is None


def test_save_single_spectrum_uses_index_fallback(workflow: MeasurementWorkflow) -> None:
    """缺少 x_axis_values 时应使用索引作为第一列。"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        pytest.skip("需要 openpyxl")

    with tempfile.TemporaryDirectory() as tmpdir:
        workflow.output_root = Path(tmpdir)
        n = 5
        workflow.context["raw_values"] = [1.0, 2.0, 3.0, 4.0, 5.0]
        workflow.context["raw_filtered_values"] = []
        workflow.context["raw_median_values"] = []
        workflow.context["fit_values"] = []
        workflow.context["x_axis_values"] = None

        path = workflow.save_single_spectrum_to_xlsx()
        assert path is not None

        wb = load_workbook(path)
        ws = wb.active
        assert ws.cell(row=2, column=1).value == 0
        assert ws.cell(row=6, column=1).value == 4


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
