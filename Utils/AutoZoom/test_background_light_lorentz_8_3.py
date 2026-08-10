"""Tests for background spectrum subtraction and Lorentzian plotting."""

import importlib.util
import inspect
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent
MODULE_NAME = "mwf_8_3_background_lorentz_test"

spec = importlib.util.spec_from_file_location(
    MODULE_NAME,
    str(PROJECT_ROOT / "0_measurement_workflow_real_virtual_same_detection_8_3.py"),
)
module = importlib.util.module_from_spec(spec)
sys.modules[MODULE_NAME] = module
spec.loader.exec_module(module)

MeasurementWorkflow = module.MeasurementWorkflow
MeasurementWorkflowGUI = module.MeasurementWorkflowGUI


class BackgroundLightLorentzTests(unittest.TestCase):
    def _make_workflow(self):
        wf = MeasurementWorkflow.__new__(MeasurementWorkflow)
        wf.log = MagicMock()
        wf.cfg = SimpleNamespace(raw_remove_above=float("inf"), median_filter_window=3)
        wf.context = {
            "x_axis_values": [500.0, 501.0, 502.0],
            "raw_values": [10.0, 12.0, 14.0],
            "raw_filtered_values": [9.0, 11.0, 13.0],
            "raw_median_values": [8.0, 10.0, 12.0],
            "fit_values": [7.0, 9.0, 11.0],
        }
        return wf

    def test_background_light_saves_fixed_xlsx_and_overwrites(self):
        wf = self._make_workflow()
        with tempfile.TemporaryDirectory() as td:
            save_dir = Path(td)
            first = MeasurementWorkflow.save_background_light_spectrum_to_xlsx(wf, save_dir=save_dir)
            self.assertEqual(first.name, "measurement_summary_background_light.xlsx")

            wf.context["raw_values"] = [1.0, 2.0, 3.0]
            second = MeasurementWorkflow.save_background_light_spectrum_to_xlsx(wf, save_dir=save_dir)
            self.assertEqual(first, second)

            wb = load_workbook(second, data_only=True)
            ws = wb.active
            self.assertEqual(ws.title, "背景光光谱数据")
            self.assertEqual(ws["B2"].value, 1.0)
            wb.close()

    def test_background_path_uses_workflow_output_root_save_date(self):
        wf = self._make_workflow()
        with tempfile.TemporaryDirectory() as td:
            wf.output_root = Path(td) / "measurement_output"
            path = MeasurementWorkflow._background_light_xlsx_path(wf)
            self.assertEqual(path.name, "measurement_summary_background_light.xlsx")
            self.assertEqual(path.parent.parent, wf.output_root / "save")

    def test_apply_background_light_correction_subtracts_raw_then_recomputes_median(self):
        wf = self._make_workflow()
        with tempfile.TemporaryDirectory() as td:
            save_dir = Path(td)
            MeasurementWorkflow.save_background_light_spectrum_to_xlsx(wf, save_dir=save_dir)

            target_path = save_dir / "measurement_summary_background_light.xlsx"
            wf._background_light_xlsx_path = lambda save_dir=None: target_path
            wf.context.update(
                {
                    "raw_values": [20.0, 23.0, 28.0],
                    "raw_filtered_values": [999.0, 999.0, 999.0],
                    "raw_median_values": [999.0, 999.0, 999.0],
                    "fit_values": [999.0, 999.0, 999.0],
                    "fit_peak": 999.0,
                }
            )
            result = {}

            MeasurementWorkflow._apply_background_light_correction(wf, result)

            self.assertEqual(wf.context["raw_values"], [10.0, 11.0, 14.0])
            self.assertEqual(wf.context["raw_filtered_values"], [10.0, 11.0, 14.0])
            self.assertEqual(wf.context["raw_median_values"], [10.5, 11.0, 12.5])
            self.assertEqual(wf.context["fit_values"], [10.5, 11.0, 12.5])
            self.assertTrue(result["background_light_applied"])
            self.assertEqual(result["fit_peak"], 12.5)

    def test_append_plot_point_records_peak_wavelength(self):
        wf = self._make_workflow()
        wf.plot_points = []
        wf.context["plot_points"] = wf.plot_points
        wf.context["fit_values"] = [1.0, 5.0, 3.0]
        wf.context["raw_median_values"] = [1.0, 4.0, 2.0]
        wf.context["fit_peak"] = 5.0

        MeasurementWorkflow._append_plot_point(wf, cycle_index=1)

        self.assertEqual(wf.plot_points[0]["fit_peak_x"], 501.0)

    def test_gui_source_contains_x_range_and_lorentz_on_median_plot(self):
        build_src = inspect.getsource(MeasurementWorkflowGUI._build_ui)
        update_src = inspect.getsource(MeasurementWorkflowGUI.update_angle_fit_peak_plot)
        self.assertIn("plot_x_min_var", build_src)
        self.assertIn("plot_x_max_var", build_src)
        self.assertIn("清空范围", build_src)
        self.assertIn("_filter_xy_by_x_range", update_src)
        self.assertIn("self._fit_lorentzian_curve(median_x, median_values)", update_src)
        self.assertIn("self.ax_median.plot(", update_src)
        self.assertIn("洛伦兹拟合", update_src)

    def test_lorentzian_model_formula(self):
        y = MeasurementWorkflowGUI._lorentzian_model([2.0], y0=1.0, z=6.0, w=2.0, c=2.0)
        self.assertAlmostEqual(float(y[0]), 4.0)

    def test_update_plot_uses_true_angle_in_list(self):
        gui = MeasurementWorkflowGUI.__new__(MeasurementWorkflowGUI)
        gui.workflow = type(
            "W",
            (),
            {
                "context": {
                    "raw_values": [1.0, 2.0],
                    "raw_median_values": [1.0, 2.0],
                    "x_axis_values": [500.0, 501.0],
                    "background_light_applied": False,
                },
                "_make_x_values_for_plot": staticmethod(lambda y, x: list(x) if x else list(range(len(y)))),
                "load_x_axis_values_from_xlsx": staticmethod(lambda: [500.0, 501.0]),
                "log": MagicMock(),
            },
        )()
        gui.fig = MagicMock()
        gui.ax_fit_peak = MagicMock()
        gui.ax_raw = MagicMock()
        gui.ax_median = MagicMock()
        gui.plot_canvas = MagicMock()
        gui.plot_status_var = MagicMock()
        gui.update_angle_fit_xy_list = MagicMock()
        gui._get_plot_x_range = lambda: (None, None)
        gui._filter_xy_by_x_range = MeasurementWorkflowGUI._filter_xy_by_x_range
        gui._fit_lorentzian_curve = lambda *args, **kwargs: None

        point = {
            "cycle_index": 1,
            "angle_deg": 25.351932505646204,
            "fit_peak": 1000.0,
            "fit_peak_x": 0.0,
            "raw_values": [1.0, 2.0],
            "raw_median_values": [1.0, 2.0],
            "x_axis_values": [500.0, 501.0],
        }
        gui.workflow.plot_points = [point]

        MeasurementWorkflowGUI.update_angle_fit_peak_plot(gui, gui.workflow)

        args, _ = gui.update_angle_fit_xy_list.call_args
        self.assertAlmostEqual(args[0][0][1], 25.351932505646204)


if __name__ == "__main__":
    unittest.main()
