from __future__ import annotations

import shutil
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT_DIR = Path(r"C:\Users\Mr\Desktop\20260905课题总结")
REF_DIR = OUT_DIR
ASSET_DIR = ROOT / "doc_assets"
ASSET_DIR.mkdir(exist_ok=True)
_FONT_PATH = r"C:\Windows\Fonts\msyh.ttc"
if Path(_FONT_PATH).exists():
    font_manager.fontManager.addfont(_FONT_PATH)
    _FONT_NAME = font_manager.FontProperties(fname=_FONT_PATH).get_name()
else:
    _FONT_NAME = "DejaVu Sans"
plt.rcParams["font.family"] = _FONT_NAME
plt.rcParams["axes.unicode_minus"] = False


def set_run_font(run, name="微软雅黑", size=10.5, bold=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", sz="6"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def format_table(table, header_fill="1F4E78"):
    table.style = "Table Grid"
    for r, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_borders(cell)
            cell.vertical_alignment = 1
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.1
                for run in p.runs:
                    set_run_font(run, size=9.5, bold=(r == 0), color=("FFFFFF" if r == 0 else "000000"))
            if r == 0:
                set_cell_shading(cell, header_fill)
            elif r % 2 == 0:
                set_cell_shading(cell, "F2F6FA")


def add_para(doc, text="", style="Normal", bold_prefix=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.line_spacing = 1.35 if style == "Normal" else 1.2
    p.paragraph_format.space_after = Pt(6 if style == "Normal" else 3)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.5 if style == "Normal" else 12, bold=style.startswith("Heading"))
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, size={1: 16, 2: 13, 3: 11.5}[level], bold=True)
    return p


def add_figure(doc, path, caption, width=6.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(6)
    r = c.add_run(caption)
    set_run_font(r, size=9.5)


def draw_flow(path, title, nodes, accent="#1F4E78"):
    fig, ax = plt.subplots(figsize=(12, 3.5), dpi=180)
    ax.set_xlim(0, len(nodes) * 2.25)
    ax.set_ylim(0, 3)
    ax.axis("off")
    ax.text(0.2, 2.72, title, fontsize=13, fontweight="bold", color="#111111", fontproperties=font_manager.FontProperties(fname=_FONT_PATH) if Path(_FONT_PATH).exists() else None)
    for i, (label, sub) in enumerate(nodes):
        x = 0.2 + i * 2.25
        box = FancyBboxPatch((x, 1.05), 1.65, 1.05, boxstyle="round,pad=0.04,rounding_size=0.08", linewidth=1.2, edgecolor=accent, facecolor="#F4F8FC")
        ax.add_patch(box)
        fp = font_manager.FontProperties(fname=_FONT_PATH) if Path(_FONT_PATH).exists() else None
        ax.text(x + 0.825, 1.67, label, ha="center", va="center", fontsize=9.5, fontweight="bold", fontproperties=fp)
        ax.text(x + 0.825, 1.30, sub, ha="center", va="center", fontsize=7.6, color="#444444", wrap=True, fontproperties=fp)
        if i < len(nodes) - 1:
            ax.add_patch(FancyArrowPatch((x + 1.68, 1.58), (x + 2.18, 1.58), arrowstyle="->", mutation_scale=12, linewidth=1.2, color=accent))
    fig.tight_layout(pad=0.6)
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def draw_layers(path):
    fig, ax = plt.subplots(figsize=(8.5, 5.0), dpi=180)
    ax.axis("off")
    layers = [
        ("应用层", "PyQt 自动聚焦控制台\n实时预览 / 曲线 / 日志 / 导出", "#EAF2F8"),
        ("闭环编排层", "AutofocusController\n触发判定 / 被动补焦 / 状态收敛", "#D9EAF7"),
        ("评分与搜索层", "FocusMetricsCalculator + FocusScorer\n14 项指标 / 参考基线 / 4 类搜索策略", "#CFE2F3"),
        ("采集与执行层", "屏幕 / 窗口 / USB / 本地视频\nNewport Picomotor Z 轴或模拟器", "#B9D7EA"),
    ]
    y = 3.8
    for i, (head, body, color) in enumerate(layers):
        rect = FancyBboxPatch((0.7, y), 7.1, 0.75, boxstyle="round,pad=0.03,rounding_size=0.05", linewidth=1, edgecolor="#1F4E78", facecolor=color)
        ax.add_patch(rect)
        fp = font_manager.FontProperties(fname=_FONT_PATH) if Path(_FONT_PATH).exists() else None
        ax.text(1.0, y + 0.46, head, fontsize=11, fontweight="bold", va="center", fontproperties=fp)
        ax.text(2.35, y + 0.36, body, fontsize=9, va="center", fontproperties=fp)
        if i < len(layers) - 1:
            ax.add_patch(FancyArrowPatch((4.25, y - 0.02), (4.25, y - 0.27), arrowstyle="->", mutation_scale=12, linewidth=1.1, color="#1F4E78"))
        y -= 1.0
    fp = font_manager.FontProperties(fname=_FONT_PATH) if Path(_FONT_PATH).exists() else None
    ax.text(4.25, 4.72, "AutoFocus 软件系统分层结构", ha="center", fontsize=13, fontweight="bold", fontproperties=fp)
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def clear_body(doc):
    body = doc._element.body
    sectPr = body.sectPr
    for child in list(body):
        if child is not sectPr:
            body.remove(child)


def setup_doc(template, title):
    doc = Document(str(template))
    clear_body(doc)
    sec = doc.sections[0]
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.9)
    sec.right_margin = Inches(0.9)
    for name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
        try:
            st = doc.styles[name]
            st.font.name = "微软雅黑"
            st._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
            st.font.color.rgb = RGBColor(0, 0, 0)
        except Exception:
            pass
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    r = p.add_run(title)
    set_run_font(r, size=16, bold=True)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(20)
    r = p2.add_run("项目文档")
    set_run_font(r, size=12, bold=True)
    return doc


def add_contents(doc, items):
    add_heading(doc, "目录", 1)
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25 if item[0] == 1 else 0.55)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(item[1])
        set_run_font(r, size=10.5)
    doc.add_page_break()


def build_detailed():
    title = "基于FocusScore多指标融合与自适应搜索的显微自动对焦系统"
    template = next(REF_DIR.glob("基于DINOv3*.docx"))
    doc = setup_doc(template, title)
    contents = [(1,"1 项目概况"),(2,"1.1 背景和基础"),(2,"1.2 场景和价值"),(2,"1.3 所需支持"),(1,"2 项目规划"),(2,"2.1 整体目标"),(2,"2.2 技术创新点"),(2,"2.3 硬件设备"),(2,"2.4 算法原理分析"),(1,"3 实施方案"),(2,"3.1 技术可行性分析"),(2,"3.2 技术细节"),(2,"3.3 系统测试与精度验证"),(1,"4 参考资料")]
    add_contents(doc, contents)

    add_heading(doc, "项目概况", 1)
    add_heading(doc, "背景和基础", 2)
    add_para(doc, "在显微观察、光谱测量和精密光学实验中，焦点位置会受到样品高度变化、载物台漂移、热变化以及操作扰动影响。传统人工调焦依赖操作者观察图像清晰度，难以持续运行，也难以用统一指标记录焦点状态。当前 AutoFocus 源码将自动对焦从测量流程中独立出来，形成了可配置的图像采集、聚焦指标计算、参考基线、Z 轴搜索和图形化操作链路。")
    add_para(doc, "项目基础来自 Focus 包与 autofocus_qt_ui 包：前者负责核心算法与 Newport Picomotor 接口，后者提供 PyQt 实时预览、ROI 框选、指标表、FocusScore 曲线、日志和离线处理工具。源码同时提供模拟 Z 轴和合成显微图像，可在无硬件条件下验证搜索逻辑；实机性能仍取决于相机、光学系统和 Newport 控制器的现场配置。")
    add_figure(doc, ASSET_DIR / "layers.png", "图1.1 AutoFocus 软件系统分层结构", width=5.9)

    add_heading(doc, "场景和价值", 2)
    add_para(doc, "系统适用于显微镜窗口或屏幕区域的连续聚焦监测，也支持 USB 相机、本地视频和模拟数据。它可以在参考图像基础上输出 FocusScore_ratio，识别聚焦退化并按策略驱动 Z 轴寻找更优位置；当不希望移动硬件时，还可启用 FocusScore 检测模式，仅保存触发图像与分数。")
    add_para(doc, "项目的直接价值是将调焦经验转化为可复现的软件流程：采集范围和 ROI 可配置，参考基线可保存，指标和分数可导出，离线视频可按时间间隔抽帧并标注。对于实验流程，系统还提供被动补焦、连续达标停止、手动暂停和模拟漂移等机制，便于做稳定性观察与参数比较。")

    add_heading(doc, "所需支持", 2)
    add_heading(doc, "硬件平台的论证和选择", 3)
    add_para(doc, "视觉采集设备可为显微镜相机、USB 相机或已有显微镜软件窗口；源码通过屏幕区域、窗口 ROI 和 USB VideoCapture 适配不同采集方式。执行设备采用 Newport 8742 Picomotor 的指定 Z 轴，控制层提供连接检查、速度/加速度配置、相对移动和关闭接口；当未接入硬件时使用 VirtualZAxis 进行模拟。")
    add_heading(doc, "软件平台的论证和选择", 3)
    add_para(doc, "系统以 Python 为主，使用 NumPy、OpenCV、Pillow 和可选的 pyautogui、pygetwindow、pywin32 完成图像与窗口采集；PyQt6/PySide2 兼容层负责桌面界面；pylablib 用于 Newport Picomotor 连接。源码还提供 pytest 测试与 CSV/图片输出，便于实验记录和回归验证。")

    add_heading(doc, "项目规划", 1)
    add_heading(doc, "整体目标", 2)
    add_para(doc, "系统以“采集—评分—判定—搜索—复测”为主闭环。采集层获得整图并裁剪 ROI；评分层同时计算整图和 ROI 的 14 项指标；参考层以多次采样均值建立基线；控制层依据 FocusScore_ratio 容差区间判断是否需要补焦；搜索层支持爬山、全扫、曲线拟合和黄金分割；执行层通过 Newport Z 轴或模拟器移动后重新测量。")
    add_figure(doc, ASSET_DIR / "focus_flow.png", "图2.1 自动对焦闭环流程", width=6.0)

    add_heading(doc, "技术创新点", 2)
    add_heading(doc, "作品难点", 3)
    for t in [
        "1. 多指标在不同样品和照明条件下的可比性：系统需要同时处理梯度、高频、边缘、亮度和颜色等指标，并通过参考基线和权重抑制单一指标波动。",
        "2. 补焦方向与搜索效率：方向探测需要在有限步数内判断上升趋势，随后复用方向完成粗搜、局部细搜和回到最佳位置，避免同一搜索周期重复判断。",
        "3. 真实硬件与离线/模拟环境兼容：窗口捕获、USB 相机、视频文件、VirtualZAxis 和 Newport Picomotor 需要共享同一控制接口，同时对依赖缺失、设备未连接和 ROI 越界给出明确错误。",
    ]:
        add_para(doc, t)
    add_heading(doc, "作品创新点", 3)
    for t in [
        "1. FocusScore_ratio 采用 14 项指标的整图/ROI 双通道计算，当前实现默认启用 highfreq_ratio、tenengrad、brenner 和 red_blue_ratio，其余指标可通过配置增加权重。",
        "2. 搜索策略可插拔：hill_climb、full_sweep、curve_fit、golden_section 和 feedback_closed_loop 由工厂函数统一创建，便于按样品焦曲线选择策略。",
        "3. 软件闭环与 UI 解耦：AutofocusController 可独立运行，PyQt 工作线程负责预览、进度、暂停、停止和日志信号，避免耗时采集阻塞界面。",
        "4. 离线数据能力：支持视频按固定间隔抽帧、基准图选择、分数标注、零分图片过滤，以及按 ROI 生成新视频。",
    ]:
        add_para(doc, t)

    add_heading(doc, "硬件设备", 2)
    p = doc.add_paragraph()
    r = p.add_run("表2.1 采集与执行平台对比")
    set_run_font(r, size=9.5)
    t = doc.add_table(rows=1, cols=4)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text, t.rows[0].cells[2].text, t.rows[0].cells[3].text = "平台", "输入/输出", "适用场景", "边界"
    for row in [("屏幕/窗口", "RGB 图像", "显微镜软件窗口", "依赖窗口权限"), ("USB 相机", "实时帧", "直接相机采集", "依赖驱动与设备"), ("本地视频", "视频帧", "离线复盘", "不执行实机移动"), ("Newport 8742", "Z 轴步进", "真实闭环补焦", "需 pylablib 与控制器")]:
        cells = t.add_row().cells
        for i, v in enumerate(row): cells[i].text = v
    format_table(t)
    add_heading(doc, "视觉采集与 ROI", 3)
    add_para(doc, "FocusMetricsCalculator 提供 screen_region、window_roi、usb_camera 和 local_video_file 等采集分支。默认截图区域为 (116, 98, 1112, 886)，默认 ROI 为 (0, 0, 300, 300)，并在计算前将 ROI 限制到实际图像边界。窗口模式优先使用 Win32 捕获，失败时回退到激活窗口后的屏幕截图。")
    add_heading(doc, "Newport Picomotor Z 轴", 3)
    add_para(doc, "ZAxisController 对 Newport 8742 连接参数、后端选择、超时、轴扫描、速度和加速度进行封装。move_relative 使用带符号步数表达方向，check_available 可在不实际移动的情况下检查 pylablib、控制器数量和连接索引。源码将 z_picomotor_conn 作为配置项，避免将设备索引硬编码在业务流程中。")
    add_heading(doc, "模拟执行设备", 3)
    add_para(doc, "FocusSimulator 由 VirtualZAxis 与 ImageGenerator 组成，可设定 peak_z、blur_scale、drift_rate、focus_degrade、pattern 和 noise_level。模拟器通过 Gaussian blur 让图像清晰度随 Z 位置变化，用于验证搜索策略和闭环停止条件，不代表真实光学系统的性能指标。")

    add_heading(doc, "算法原理分析", 2)
    p = doc.add_paragraph()
    r = p.add_run("表2.2 闭环搜索策略对比")
    set_run_font(r, size=9.5)
    t = doc.add_table(rows=1, cols=4)
    for i, v in enumerate(("策略", "核心方法", "优点", "适用条件")): t.rows[0].cells[i].text = v
    for row in [("hill_climb", "方向探测+爬坡+局部细搜", "移动量较小，支持方向复用", "焦点峰值较连续"), ("full_sweep", "固定范围逐点扫描", "不依赖先验曲线", "首次标定或峰值未知"), ("curve_fit", "离散采样+抛物线拟合", "可预测峰值位置", "焦曲线近似单峰"), ("golden_section", "区间收缩", "采样效率高", "峰值区间已知")]:
        cells = t.add_row().cells
        for i, v in enumerate(row): cells[i].text = v
    format_table(t)
    add_heading(doc, "聚焦指标与参考基线", 3)
    add_para(doc, "FocusMetricsCalculator 对 RGB 图像计算 tenengrad、laplacian_var、brenner、highfreq_ratio、local_contrast、edge_width、halo_width、brightness_mean、brightness_std、red_blue_ratio、modified_laplacian、dct_energy、smd 和 entropy 共 14 项指标。FocusScorer 先以多次采样建立参考均值，再对当前 ROI 的各项指标计算 current/reference，并按权重求加权均值；单项比值限制在 [0, 2.5]，防止异常值支配总分。")
    add_figure(doc, ASSET_DIR / "metrics.png", "图2.2 聚焦指标计算与评分组成", width=6.0)
    add_heading(doc, "闭环触发与搜索策略", 3)
    add_para(doc, "默认触发阈值为 0.95，绝对值模式下允许区间为 [0.95, 1.05]；连续 3 轮超出区间后触发补焦。hill_climb 先进行多点方向探测，再沿上升方向迭代，最后回到最佳位置并做局部细搜；full_sweep 进行固定范围扫描；curve_fit 以抛物线拟合峰值后微调；golden_section 在 bracket 区间内递归收敛。搜索停止还受最大迭代次数、耐心轮数、总步数和手动停止状态约束。")

    add_heading(doc, "实施方案", 1)
    add_heading(doc, "技术可行性分析", 2)
    add_para(doc, "从源码结构看，核心模块边界清晰：config.py 保存参数，metrics.py 负责图像与采集，scorer.py 负责参考与评分，search.py 负责策略，controller.py 负责一次完整补焦，z_axis.py 负责硬件，simulator.py 提供无硬件替代，main.py 和 PyQt UI 提供运行入口。因此可以先用模拟器和离线视频完成算法验证，再接入真实窗口、相机和 Z 轴。")
    add_para(doc, "风险主要来自真实采集环境：ROI 内容变化、曝光/照明变化、参考图不稳定会使 FocusScore_ratio 偏离纯粹的清晰度含义；pylablib、USB 驱动或窗口权限缺失会阻止实机移动。部署时应先运行 check_available、建立参考基线，并设置被动模式的最大尝试次数与机械行程保护。")
    add_heading(doc, "技术细节", 2)
    add_heading(doc, "系统主程序设计", 3)
    add_para(doc, "命令行入口支持 --build-ref、--cycles、--interval、--demo-sim、--strategy、--roi 等参数。每轮由 check_and_autofocus() 完成采集、指标计算、触发判定、可选闭环搜索和结果保存；main.py 同时写入日志与 CSV，便于后续分析。")
    add_heading(doc, "PyQt 交互界面设计", 3)
    add_para(doc, "autofocus_qt_ui 提供采集模式、设备选择、截图区域、ROI、搜索策略、Z 轴参数、触发阈值、被动补焦和 FocusScore 检测开关。右侧区域显示实时预览、14 项指标表、趋势曲线、横竖剖面和运行日志；后台 QThread 负责闭环、离线检测和视频 ROI 裁剪。")
    add_heading(doc, "离线处理设计", 3)
    add_para(doc, "OfflineDatasetDetector 可从视频按默认 3 秒间隔抽帧，也可读取图片文件夹；选择基准图后，对每张图计算 FocusScore 并在图像上方标注。VideoRoiCropper 对视频逐帧裁剪 ROI，自动限制越界坐标，保持源视频帧率，并按 mp4v、XVID、MJPG 顺序尝试编码。")
    add_heading(doc, "系统测试与精度验证", 2)
    add_para(doc, "本次源码回归测试执行 `python -m pytest -q`，结果为 11 passed，覆盖 ROI 边界限制、输出路径、单帧裁剪、视频信息读取、完整视频裁剪、进度回调、后台线程成功与异常信号等路径。该结果证明离线视频 ROI 模块和 Qt worker 的软件行为满足现有测试断言。")
    add_para(doc, "当前仓库未提供本轮真实显微镜图像、实机 Newport 运动日志或统一的光学重复定位测试数据，因此本文不把模拟峰值、FocusScore_ratio 或软件单元测试结果表述为实机对焦精度。后续现场验证应补充不同样品、不同照明、正负方向扰动、长时间漂移和机械限位场景。")
    add_figure(doc, ASSET_DIR / "test_pipeline.png", "图3.1 软件测试与现场验证边界", width=6.0)

    add_heading(doc, "参考资料", 1)
    refs = [
        "[1] Bradski G. The OpenCV Library. Dr. Dobb's Journal of Software Tools, 2000.",
        "[2] Gonzalez R C, Woods R E. Digital Image Processing. Pearson, 2018.",
        "[3] Krotkov E. Focusing. International Journal of Computer Vision, 1988, 1: 223-237.",
        "[4] Newport Corporation. Picomotor Controller 8742 User's Manual, 2024.",
        "[5] Riverbank Computing. PyQt Documentation. https://www.riverbankcomputing.com/software/pyqt/",
        "[6] Python Software Foundation. Python 3 Documentation. https://docs.python.org/3/",
    ]
    for ref in refs: add_para(doc, ref)
    out = OUT_DIR / (title + ".docx")
    doc.save(str(out))
    return out


def build_concise():
    title = "基于多源图像采集与Newport Z轴闭环伺服的智能稳焦系统"
    template = next(REF_DIR.glob("基于UCC*.docx"))
    doc = setup_doc(template, title)
    add_heading(doc, "项目概况", 1)
    add_heading(doc, "背景和基础", 2)
    add_para(doc, "显微观察和光谱测量过程中，样品高度、载物台漂移及环境变化会造成焦点偏移。AutoFocus 项目将图像清晰度度量、参考基线和 Z 轴运动封装为统一闭环，并提供桌面控制台和无硬件模拟模式。系统的目标是持续监测 ROI 内的聚焦状态，在分数超出允许区间时自动搜索更优 Z 位置，并把每轮图像、指标、分数和日志保存下来。")
    add_para(doc, "项目已完成核心 Focus 包、PyQt UI、离线视频评分与 ROI 视频裁剪模块的代码实现；真实硬件闭环的效果需要在具体显微镜、相机和 Newport 8742 控制器上进一步标定。")
    add_figure(doc, ASSET_DIR / "focus_flow.png", "图1.1 自动稳焦闭环架构", width=6.0)
    add_heading(doc, "所需支持", 2)
    add_heading(doc, "硬件平台的论证和选择", 3)
    add_para(doc, "前端可接入显微镜软件窗口、屏幕区域或 USB 相机；本地视频用于离线复盘。执行端采用 Newport 8742 Picomotor 的一个 Z 轴，支持相对步进、速度/加速度设置、设备检查与安全关闭；没有控制器时以 VirtualZAxis 模拟移动。")
    add_heading(doc, "软件平台的论证和选择", 3)
    add_para(doc, "软件采用 Python、NumPy、OpenCV、Pillow、PyQt 兼容层和可选 pylablib。核心接口通过 AutofocusConfig 统一参数，通过 FocusMetricsCalculator、FocusScorer、AutofocusController 和 ZAxisController 分层实现采集、评分、决策和执行。")

    add_heading(doc, "项目规划", 1)
    add_heading(doc, "整体目标", 2)
    add_para(doc, "系统遵循“感知—决策—执行—复测”架构。感知层采集整图并裁剪 ROI；决策层计算 14 项指标并生成 FocusScore_ratio；执行层依据触发状态选择 hill_climb、full_sweep、curve_fit 或 golden_section；应用层提供参考建立、实时曲线、日志、暂停/停止和结果导出。")
    add_heading(doc, "技术创新点", 2)
    add_heading(doc, "作品难点", 3)
    for t in [
        "1. 参考基线与样品变化的耦合：FocusScore_ratio 是相对指标，ROI 内容、曝光和照明变化都会影响结果，需要把清晰度判断与样品稳定性一起管理。",
        "2. 方向判断和局部细搜的衔接：一次搜索周期内标记方向并持续复用，避免重复探测造成额外机械移动。",
        "3. UI、离线任务和硬件控制并发：视频抽帧、评分与 Z 轴移动不能阻塞主界面，后台线程需要传递进度、日志和完成状态。",
    ]: add_para(doc, t)
    add_heading(doc, "作品创新点", 3)
    for t in [
        "1. 14 项指标的整图/ROI 双列显示与加权 FocusScore_ratio，默认权重集中于高频和梯度指标，其他指标可按实验启用。",
        "2. 同一控制器兼容真实 Newport Picomotor、模拟 Z 轴和离线视频，支持从算法调试平滑迁移到设备联调。",
        "3. 将实时闭环、FocusScore 检测、离线视频评分和 ROI 裁剪放入一个 PyQt 工作台，形成可回溯的实验数据链路。",
    ]: add_para(doc, t)

    add_heading(doc, "主要硬件设备", 2)
    add_heading(doc, "图像采集与聚焦 ROI", 3)
    add_para(doc, "采集模块支持 screen_region、window_roi、usb_camera 和 local_video_file。窗口捕获优先调用 Win32 API，必要时回退到窗口置顶后的屏幕截图；USB 模式通过 OpenCV VideoCapture 读取帧。ROI 在计算前执行边界钳制，确保宽高为正。")
    add_heading(doc, "Newport 8742 Z轴闭环伺服", 3)
    add_para(doc, "ZAxisController 负责检测 Picomotor 数量、选择连接索引和后端、设置速度/加速度并发送带符号步进。AutofocusController 在每轮评分低于阈值且连续达到触发次数后调用搜索策略，搜索完成后回到最佳位置并复测。")
    add_heading(doc, "算法原理分析", 2)
    add_heading(doc, "FocusScore多指标融合", 3)
    add_para(doc, "系统计算 tenengrad、laplacian_var、brenner、highfreq_ratio、local_contrast、edge_width、halo_width、brightness_mean、brightness_std、red_blue_ratio、modified_laplacian、dct_energy、smd 和 entropy。FocusScorer 以参考图多次采样均值为分母，计算加权 current/reference；默认触发区间为 [0.95, 1.05]。")
    add_heading(doc, "多策略闭环搜索", 3)
    add_para(doc, "hill_climb 适合焦点峰值附近快速爬坡，full_sweep 适合焦曲线未知的初次扫描，curve_fit 通过离散采样预测峰值，golden_section 用区间收缩细化峰值。被动补焦模式可在分数未恢复时连续尝试，并通过最大尝试次数和连续达标次数限制运行范围。")

    add_heading(doc, "实施方案", 1)
    add_heading(doc, "测试环境与平台", 2)
    add_para(doc, "软件测试可在无硬件环境下运行，使用 FocusSimulator 生成不同峰值和漂移条件；离线测试使用 OpenCV 生成的短视频或用户提供的视频/图片文件夹。真实部署需补充显微镜窗口权限、USB 驱动、pylablib 和 Newport 控制器连接检查。")
    add_heading(doc, "离线处理与交互测试", 2)
    add_para(doc, "现有 pytest 回归结果为 11 passed，覆盖 ROI 越界钳制、视频信息读取、逐帧裁剪、编码输出、进度回调及 worker 成功/失败信号。该证据属于软件行为验证；尚未形成真实光学焦点重复定位、长时间漂移或 Z 轴机械误差的现场测量报告。")
    add_heading(doc, "运行稳定性与边界条件", 2)
    add_para(doc, "系统对以下异常提供处理路径：未建立参考时禁止评分、窗口不存在、视频无法打开、ROI 宽高无效、ROI 越界、pylablib 缺失、控制器未检测到以及后台任务异常。部署时建议先以 detection_only 模式观察分数，再开启 Z 轴移动，并保留日志与触发图像。")
    add_figure(doc, ASSET_DIR / "test_pipeline.png", "图3.1 软件回归、离线验证与实机验证边界", width=6.0)
    add_heading(doc, "参考资料", 1)
    for ref in [
        "[1] Bradski G. The OpenCV Library. Dr. Dobb's Journal of Software Tools, 2000.",
        "[2] Krotkov E. Focusing. International Journal of Computer Vision, 1988, 1: 223-237.",
        "[3] Newport Corporation. Picomotor Controller 8742 User's Manual, 2024.",
        "[4] Python Software Foundation. Python 3 Documentation. https://docs.python.org/3/",
        "[5] Qt Project. Qt for Python Documentation. https://doc.qt.io/qtforpython/",
    ]: add_para(doc, ref)
    out = OUT_DIR / (title + ".docx")
    doc.save(str(out))
    return out


def main():
    draw_layers(ASSET_DIR / "layers.png")
    draw_flow(ASSET_DIR / "focus_flow.png", "闭环自动对焦主流程", [("图像采集", "屏幕 / 窗口 / USB / 视频"), ("ROI指标", "14项指标"), ("参考评分", "FocusScore_ratio"), ("触发判定", "阈值 + 连续计数"), ("Z轴搜索", "4类策略"), ("复测记录", "图像 / CSV / 日志")])
    draw_flow(ASSET_DIR / "metrics.png", "指标计算与评分组成", [("RGB图像", "整图与ROI"), ("梯度与频域", "Tenengrad / FFT / DCT"), ("边缘与统计", "边缘 / 亮度 / 熵"), ("参考基线", "多次采样均值"), ("加权融合", "current / reference"), ("输出分数", "触发与趋势")], accent="#2E7D32")
    draw_flow(ASSET_DIR / "test_pipeline.png", "验证证据边界", [("pytest", "11 passed"), ("模拟器", "算法路径"), ("离线视频", "抽帧 / ROI"), ("真实相机", "待现场"), ("Newport Z轴", "待联调"), ("光学指标", "待标定")], accent="#8A4B08")
    outputs = [build_detailed(), build_concise()]
    for p in outputs: print(p)


if __name__ == "__main__":
    main()
